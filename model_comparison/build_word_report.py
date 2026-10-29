"""Word report: Sections 3-5 (methodology, results, contributions).

Reads saved outputs only (run the model scripts, compare_models.py,
score_variant.py and threshold_event_metrics.py first). Figures are written
to model_comparison/report_figures/.
"""

import json
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_report as br  # noqa: E402  (loads all saved outputs + plot style)

ROOT = br.ROOT
CMP = br.CMP
FIG = ROOT / "model_comparison/report_figures"
FIG.mkdir(exist_ok=True)
OUT = ROOT / "model_comparison/regime_report_sections_3_5.docx"
TARGET = br.TARGET
MODELS, PHASES = br.MODELS, br.PHASES
LOGIT_C, RF_C, INK2, MUTED = br.LOGIT_C, br.RF_C, br.INK2, br.MUTED
STRESS_C, NORMAL_C = "#6250d6", "#1baf7a"

panel = pd.read_csv(ROOT / "logit/outputs/logit_decomposition/model_panel.csv", parse_dates=["date", "event_date"])
thr = pd.read_csv(CMP / "threshold_metrics.csv")
ev = pd.read_csv(CMP / "event_level_detail.csv", parse_dates=["onset_event_date"])
ev_sum = pd.read_csv(CMP / "event_level_summary.csv")
score_res = pd.read_csv(ROOT / "model_comparison/outputs_score/score_vs_flag.csv")


# ============================================================ figures

def save(fig, name):
    path = FIG / name
    fig.savefig(path, dpi=200, bbox_inches="tight")
    plt.close(fig)
    return path


def fig_design():
    fig, ax = plt.subplots(figsize=(9, 3.4))
    rows = []
    for m in ["2022-07", "2022-08", "2022-09", "2022-10", "2022-11"]:
        start = pd.Timestamp(m + "-01")
        rows.append((f"Validate {start:%b}", "2022-04-01", str(start.date()), str(start.date()),
                     str((start + pd.offsets.MonthBegin(1)).date())))
    rows.append(("Final test", "2022-04-01", "2022-11-02", "2023-02-01", "2023-04-30"))
    for i, (label, a, b, c, d) in enumerate(rows):
        y = len(rows) - i
        ax.barh(y, pd.Timestamp(b) - pd.Timestamp(a), left=pd.Timestamp(a), color=LOGIT_C, alpha=.35, height=.55)
        ax.barh(y, pd.Timestamp(d) - pd.Timestamp(c), left=pd.Timestamp(c), color=RF_C, height=.55)
        ax.text(pd.Timestamp("2022-03-25"), y, label, ha="right", va="center", fontsize=8.5)
    ax.axvspan(pd.Timestamp("2022-12-01"), pd.Timestamp("2023-01-31"), color="#f1f0ec")
    ax.text(pd.Timestamp("2022-12-31"), 3.5, "no on-chain\ndata", ha="center", fontsize=7.5, color=INK2)
    ax.set_yticks([])
    ax.set_xlim(pd.Timestamp("2022-01-15"), pd.Timestamp("2023-05-10"))
    ax.grid(axis="y", visible=False)
    ax.barh(0, 0, color=LOGIT_C, alpha=.35, label="Training data (outcomes known)")
    ax.barh(0, 0, color=RF_C, label="Forecast period (never seen in training)")
    ax.legend(loc="lower right", fontsize=8, ncol=2, bbox_to_anchor=(1, -0.32))
    ax.set_title("Chronological evaluation: expanding-window validation in 2022, one frozen test on 2023")
    return save(fig, "fig01_design.png")


def fig_context():
    fig, axes = plt.subplots(1, 2, figsize=(10, 3.6), gridspec_kw={"width_ratios": [215, 89]}, sharey=True)
    coins = ["DAI", "USDC", "USDT"]
    for ax, yr in zip(axes, [2022, 2023]):
        d = panel[panel.event_date.dt.year.eq(yr) if yr == 2023 else panel.date.dt.year.eq(2022)]
        days = d.drop_duplicates("date")
        for day in days[days.stress.eq(1)].event_date:
            ax.axvspan(day - pd.Timedelta(hours=12), day + pd.Timedelta(hours=12), color="#e7e3f8", lw=0)
        for k, c in enumerate(coins):
            g = d[d.coin.eq(c) & d[TARGET].eq(1)]
            ax.scatter(g.event_date, np.full(len(g), k), marker="|", s=90, color=br.INK, lw=1.6)
        ax.set_yticks(range(3), coins)
        ax.set_ylim(-0.6, 2.6)
        ax.grid(axis="y", visible=False)
        ax.set_title(f"{yr}: stressed days {days.stress.mean():.0%}, breach days {int(d[TARGET].sum())}", fontsize=9)
        for lbl in ax.get_xticklabels():
            lbl.set_rotation(30)
            lbl.set_ha("right")
    axes[0].text(0.01, -0.33, "Purple shading = regime classifier says 'stressed'.  Black ticks = depeg breach days.",
                 transform=axes[0].transAxes, fontsize=8, color=INK2)
    fig.tight_layout()
    return save(fig, "fig02_regime_and_breaches.png")


def fig_ap():
    fig, axes = plt.subplots(1, 2, figsize=(8.5, 3.8))
    for ax, phase in zip(axes, PHASES):
        x = np.arange(2)
        for k, model in enumerate(MODELS):
            vals = [br.rm(model, phase, s) for s in ["M1", "M2"]]
            bars = ax.bar(x + (k - .5) * .36, vals, width=.34, color=br.MODEL_C[model], edgecolor="white",
                          linewidth=2, label=model)
            for b, v in zip(bars, vals):
                ax.text(b.get_x() + b.get_width() / 2, v + .015, f"{v:.2f}", ha="center", fontsize=8)
        base = br.preds[("Logit", phase)][TARGET].mean()
        ax.axhline(base, color=MUTED, ls=":", lw=1)
        ax.text(1.5, base + .015, f"no-skill = {base:.2f}", fontsize=7, color=INK2, ha="right")
        ax.set_xticks(x, ["M1\non-chain only", "M2\n+ regime flag"])
        ax.set_ylim(0, 1)
        ax.set_title(phase)
        ax.grid(axis="x", visible=False)
    axes[0].set_ylabel("AP (average precision)")
    axes[0].legend(fontsize=8, loc="upper left")
    fig.tight_layout()
    return save(fig, "fig03_pr_auc.png")


def fig_placebo():
    fig, axes = plt.subplots(2, 2, figsize=(9, 6.2))
    for i, phase in enumerate(PHASES):
        for j, model in enumerate(MODELS):
            ax = axes[j, i]
            g = br.plac[(br.plac.model == model) & (br.plac.phase == phase)]["delta_average_precision"].dropna()
            real = br.rm(model, phase, "M2") - br.rm(model, phase, "M1")
            ax.hist(g, bins=20, color="#c9c8c3", edgecolor="white")
            ax.axvline(real, color=br.MODEL_C[model], lw=2.5)
            ax.axvline(0, color=MUTED, lw=.8, ls=":")
            p = br.ps(model, phase, "average_precision").placebo_p_value
            ax.text(.98, .95, f"real flag: {real:+.3f}\n{p:.0%} of fake flags\ndo as well",
                    transform=ax.transAxes, ha="right", va="top", fontsize=8)
            ax.set_title(f"{model}: {phase}", fontsize=9)
            ax.set_xlabel("AP gain from adding the flag (M2 - M1)")
            ax.set_ylabel("Number of fake flags")
    fig.tight_layout()
    return save(fig, "fig04_placebo.png")


def fig_svb():
    fig, axes = plt.subplots(2, 1, figsize=(9, 6.6), sharex=True)
    for ax, coin in zip(axes, ["USDC", "DAI"]):
        for model in MODELS:
            p = br.preds[(model, "2023 frozen test")]
            g = p[p.coin.eq(coin) & p.event_date.between("2023-03-01", "2023-03-20")].sort_values("event_date")
            ax.plot(g.event_date, g.p_m1_oos, color=br.MODEL_C[model], ls="--", lw=1.4, label=f"{model} M1")
            ax.plot(g.event_date, g.p_m2_oos, color=br.MODEL_C[model], marker="o", ms=3.5, label=f"{model} M2")
        hit = g[g[TARGET].eq(1)]
        ax.scatter(hit.event_date, np.full(len(hit), 1.0), marker="v", color=br.INK, s=30, label="Actual breach")
        for d in g[g.stress.eq(1)].event_date:
            ax.axvspan(d - pd.Timedelta(hours=12), d + pd.Timedelta(hours=12), color="#e7e3f8", zorder=0)
        ax.set_ylim(0, 1.05)
        ax.set_ylabel("Forecast probability")
        ax.set_title(f"{coin} around the SVB collapse (purple = flag 'stressed')", fontsize=9.5)
    axes[0].legend(fontsize=7.5, ncol=3, loc="upper left")
    fig.autofmt_xdate()
    fig.tight_layout()
    return save(fig, "fig05_svb_paths.png")


def fig_contrast():
    fig, ax = plt.subplots(figsize=(8.5, 3.4))
    for model in MODELS:
        p = br.preds[(model, "2023 frozen test")]
        g = p[p.coin.isin(["USDC", "DAI"]) & p.event_date.between("2023-03-01", "2023-03-20")]
        g = g.groupby("event_date").m2_stress_contrast.mean()
        ax.plot(g.index, 100 * g.values, color=br.MODEL_C[model], marker="o", ms=3.5, label=model)
    ax.axhline(0, color=MUTED, lw=.8)
    ax.set_ylabel("Percentage points")
    ax.set_title("How much a 'stressed' flag would raise the M2 forecast (USDC/DAI average)", fontsize=9.5)
    ax.legend(fontsize=8)
    fig.autofmt_xdate()
    fig.tight_layout()
    return save(fig, "fig06_flag_contrast.png")


def fig_slopes():
    s = br.slopes.copy()
    s["f"] = s.feature.str.replace("log1p_", "").str.replace("_", " ")
    fig, ax = plt.subplots(figsize=(8, 4))
    y = np.arange(len(s))
    for off, col, se, c, lab in [(-.15, "normal_log_odds_slope", "normal_clustered_se", NORMAL_C, "Normal market"),
                                 (.15, "stressed_log_odds_slope", "stressed_clustered_se", STRESS_C, "Stressed market")]:
        ax.errorbar(s[col], y + off, xerr=1.96 * s[se], fmt="o", color=c, ms=5, capsize=2, label=lab)
    ax.axvline(0, color=MUTED, lw=.8)
    ax.set_yticks(y, s.f)
    ax.invert_yaxis()
    ax.set_xlabel("Effect on log-odds of a depeg (per 1 SD of the feature), with 95% CI")
    ax.legend(fontsize=8, loc="lower right")
    ax.grid(axis="y", visible=False)
    ax.set_title("Logit M3: the same on-chain feature, two regimes", fontsize=10)
    fig.tight_layout()
    return save(fig, "fig07_logit_slopes.png")


def fig_pdp():
    feats = list(dict.fromkeys(br.rf_pdp.feature))
    fig, axes = plt.subplots(2, 4, figsize=(11, 5.2))
    for ax, f in zip(axes.flat, feats):
        g = br.rf_pdp[br.rf_pdp.feature.eq(f)]
        ax.plot(g.grid_value, g.pd_normal, color=NORMAL_C, label="Normal")
        ax.plot(g.grid_value, g.pd_stressed, color=STRESS_C, label="Stressed")
        ax.set_title(f.replace("log1p_", "").replace("_", " "), fontsize=9)
        ax.tick_params(labelsize=7)
    axes.flat[0].legend(fontsize=7)
    axes.flat[0].set_ylabel("Avg. depeg probability", fontsize=8)
    axes.flat[4].set_ylabel("Avg. depeg probability", fontsize=8)
    ax = axes.flat[-1]
    imp = br.rf_int_imp.sort_values("mdi_importance")
    ax.barh(imp.feature.str.replace("log1p_", "").str.replace("_", " "), imp.mdi_importance, color=RF_C)
    ax.set_title("Importance in the forest", fontsize=9)
    ax.tick_params(labelsize=7)
    ax.grid(axis="y", visible=False)
    fig.text(.5, -.01, "x-axis: feature value (log scale)", ha="center", fontsize=8, color=INK2)
    fig.tight_layout()
    return save(fig, "fig08_rf_partial_dependence.png")


def fig_coef_imp():
    fig, axes = plt.subplots(1, 2, figsize=(9.5, 3.8))
    c = br.logit_coef[(br.logit_coef.fold == "frozen_2022") & (br.logit_coef.feature != "intercept")]
    r = br.rf_imp[br.rf_imp.fold == "frozen_2022"]
    order = list(c[c.model == "M2"].feature)
    labels = [f.replace("log1p_", "").replace("_", " ") for f in order]
    for k, spec in enumerate(["M1", "M2"]):
        cc = c[c.model == spec].set_index("feature").coefficient.reindex(order)
        rr = r[r.model == spec].set_index("feature").mdi_importance.reindex(order)
        y = np.arange(len(order)) + (k - .5) * .38
        axes[0].barh(y, cc.values, height=.36, color=LOGIT_C, alpha=.45 + .55 * k, label=spec)
        axes[1].barh(y, rr.values, height=.36, color=RF_C, alpha=.45 + .55 * k, label=spec)
    for ax, t in zip(axes, ["Logit coefficient (per 1 SD)", "Random forest importance"]):
        ax.set_yticks(np.arange(len(order)), labels, fontsize=8)
        ax.invert_yaxis()
        ax.set_title(t, fontsize=9.5)
        ax.legend(fontsize=7)
        ax.grid(axis="y", visible=False)
    axes[0].axvline(0, color=MUTED, lw=.8)
    fig.tight_layout()
    return save(fig, "fig09_coefficients_importance.png")


# ============================================================ docx helpers

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
for s in ["Heading 1", "Heading 2", "Heading 3"]:
    doc.styles[s].font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
counters = {"table": 0, "figure": 0, "prefix": ""}


def para(text, bold_lead=None, italic=False, size=None, after=6):
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(after)
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def table(df, caption, note=None, font=9, highlight=None):
    counters["table"] += 1
    cap = doc.add_paragraph()
    cap.add_run(f"Table {counters['prefix']}{counters['table']}. ").bold = True
    cap.add_run(caption).bold = True
    cap.paragraph_format.space_after = Pt(2)
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, col in enumerate(df.columns):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(col))
        run.bold = True
        run.font.size = Pt(font)
        shade(cell, "DCE6F2")
    for ri, row in enumerate(df.itertuples(index=False)):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run("" if (isinstance(v, float) and np.isnan(v)) else str(v))
            run.font.size = Pt(font)
            if i > 0:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if highlight is not None and highlight(ri, row):
                shade(cells[i], "FFF4E0")
    if note:
        p = doc.add_paragraph()
        r = p.add_run(note)
        r.italic = True
        r.font.size = Pt(8.5)
        p.paragraph_format.space_after = Pt(10)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def figure(path, caption, note=None, width=6.3):
    counters["figure"] += 1
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.add_run(f"Figure {counters['prefix']}{counters['figure']}. ").bold = True
    p.add_run(caption)
    if note:
        r = p.add_run(" " + note)
        r.italic = True
        r.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(10)


def f3(x):
    return "" if pd.isna(x) else f"{x:.3f}"


def sgn(x, d=3):
    return "" if pd.isna(x) else f"{x:+.{d}f}"


def d_(model, phase, metric="average_precision"):
    return br.rm(model, phase, "M2", col=metric) - br.rm(model, phase, "M1", col=metric)


def callout(title, lines):
    """Shaded one-cell box for key explanations."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    shade(cell, "EEF3FA")
    cell.text = ""
    r = cell.paragraphs[0].add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    for line in lines:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        if isinstance(line, tuple):
            p.add_run(line[0]).bold = True
            line = line[1]
        run = p.add_run(line)
        run.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ============================================================ numbers used in text

wald, rf_int = br.wald, br.rf_int
p22 = panel[panel.date.dt.year.eq(2022)]
p23 = panel[panel.date.dt.year.eq(2023)]
br22 = p22.groupby("stress")[TARGET].agg(["size", "sum", "mean"])
br23 = p23.groupby("stress")[TARGET].agg(["size", "sum", "mean"])
n_fake_logit22 = int(round(br.ps("Logit", PHASES[0], "average_precision").placebo_p_value * 90)) - 1

figs = {
    "design": fig_design(), "context": fig_context(), "ap": fig_ap(), "placebo": fig_placebo(),
    "svb": fig_svb(), "contrast": fig_contrast(), "slopes": fig_slopes(), "pdp": fig_pdp(),
    "coef": fig_coef_imp(),
}

# ============================================================ title
doc.add_heading("Regime-Aware Early Warning for Stablecoin Depegs", level=0)
para("Sections 3-5: Methodology, Results, and Contributions", italic=True)

# ============================================================ 3 METHODOLOGY
doc.add_heading("3. Methodology and Empirical Framework", level=1)

doc.add_heading("3.1 Approach", level=2)
para("Our question has two parts. (1) Prediction: does adding the market regime make on-chain models forecast "
     "depegs better? We judge this only on data the models have never seen. (2) Relationships: does the same "
     "on-chain signal carry different risk in stressed and calm markets? We answer this by looking inside the "
     "fitted models. For both, we compare pairs of models that are identical except for the regime "
     "information, so any difference can be attributed to it.")

doc.add_heading("3.2 Outcome, predictors and regime", level=2)
bullet(" Whether a coin (DAI, USDC, USDT) breaches its Phase 1 dynamic depeg band on day t+1. The band "
       "narrows for more heavily traded coins and uses only the previous month's volume, so it is known in "
       "advance. Note: the label marks every breach day, including days inside an ongoing episode, rather than "
       "only new breaches as planned. Section 4.5 therefore also evaluates the first day of each new episode.",
       "Outcome.")
bullet(" Seven on-chain measures for day t (transfers, volume, unique senders and receivers, minted, burned, "
       "large-transfer volume), log-transformed.", "Predictors.")
bullet(" Our Phase 1 rule: a day is 'stressed' when at least two of three signals agree: bearish BTC, bearish "
       "ETH, and Fear & Greed below 40. It was fixed before testing and is a simple proxy for crypto-wide "
       "stress.", "Regime.")

doc.add_heading("3.3 Models", level=2)
table(pd.DataFrame([
    ["M1 (baseline)", "7 on-chain features", "How well does on-chain activity alone predict depegs?"],
    ["M2 (regime-aware)", "M1 + stress flag (0/1)", "Does the regime improve forecasts? (Part 1)"],
    ["M3 (logit only)", "M2 + each feature x stress", "Do feature effects change under stress? (Part 2)"],
], columns=["Specification", "Inputs", "Question"]), "Model specifications")
bullet(" Our main benchmark. Each feature pushes risk in one direction, so results are easy to read; in M2 "
       "the flag raises or lowers baseline risk on stressed days.", "Logistic regression (logit).")
bullet(" A flexible check: 500 decision trees that learn 'if-then' rules, including thresholds and interactions "
       "we do not specify. If the regime matters in a non-linear way, the forest should find it.",
       "Random forest (RF).")
para("Settings for both models were fixed in advance rather than tuned. With only 8 breaches in the 2022 "
     "validation months, tuning would mostly fit noise.")

doc.add_heading("3.4 Evaluation design", level=2)
para("Everything runs forward in time (Figure 1). In 2022, each month from July to November is forecast by a "
     "model trained only on earlier data. A final model is then trained on all of 2022, frozen, and used once "
     "to forecast February-April 2023. No 2023 outcome is used for anything. 2022 is dominated by the Terra "
     "collapse and 2023 by SVB, so the test asks whether lessons from one crisis carry over to a different one.")
figure(figs["design"], "Forward-in-time evaluation design.")
table(pd.DataFrame([
    ["Training (2022, Apr-Nov)", 645, int(p22[TARGET].sum()), f"{p22.stress.mean():.0%}"],
    ["2022 validation forecasts (Jul-Nov)", 372, "8 (all in July)",
     f"{br.preds[('Logit', PHASES[0])].stress.mean():.0%}"],
    ["2023 test forecasts (Feb-Apr)", 267, int(p23[TARGET].sum()), f"{p23.stress.mean():.0%} (3 days)"],
], columns=["Sample", "Coin-days", "Breach days", "Stressed days"]), "Sample sizes")
figure(figs["context"], "When the regime flag is on (purple) and when depegs happen (black).",
       "In 2022 breaches cluster in stressed periods; in 2023 the flag is on for only three days.")

doc.add_heading("3.5 How we judge 'adds value'", level=2)
para("Depegs are rare, so accuracy is misleading (always predicting 'no depeg' is about 98% accurate). We use:")
bullet(" How well the model ranks breach days above calm days, focusing on the rare breaches. It summarises "
       "the precision-recall curve (our PR-AUC measure). A no-skill model scores the breach rate; 1 is perfect.",
       "AP, average precision (main metric).")
bullet(" Probability that a random breach day is ranked above a random calm day (0.5 = coin flip).", "ROC-AUC.")
bullet(" Average squared gap between forecast probability and outcome; lower is better.", "Brier score.")
bullet(" Precision, recall, false alarms and missed breaches using an alert cut-off fixed on 2022, plus "
       "whether each new 2023 episode was warned of in advance.", "Alerts and early warning.")
para("Two checks separate real gains from luck. A block bootstrap resamples forecast weeks 2,000 times to give "
     "a 95% range for M2 minus M1. A placebo test asks whether the real flag beats fake ones:")
callout("How to read the placebo p-value", [
    "Does the real flag help because of when it switches on, or would any on/off variable with long "
    "stressed spells help just as much? To check, we create 89 fake flags by sliding the real flag's history "
    "along the calendar by 30 to 646 days. Each fake has the same number and length of stressed spells, but at "
    "the wrong dates. We rerun the whole analysis with each fake.",
    ("Placebo p = ", "(1 + number of fakes that improve the model at least as much as the real flag) / (1 + 89)."),
    ("Example: ", f"for the logit in 2022, {n_fake_logit22} of the 89 fakes improved AP as much as the real flag, "
                  f"so p = {n_fake_logit22 + 1}/90 = {(n_fake_logit22 + 1) / 90:.2f}. The real flag is not special. "
                  "If no fake did as well, p would be 1/90 = 0.01, meaning the real timing matters."),
    ("Rule of thumb: ", "p of 0.05 or less is convincing; p above 0.10 means a random flag with the same "
                        "pattern does about as well."),
])
para("All results are predictive associations, not causal effects. We assume day-t data are available before "
     "forecasting day t+1, and check this by also using the flag from earlier days (Section 4.6).")

# ============================================================ 4 RESULTS
doc.add_heading("4. Results and Interpretation", level=1)

doc.add_heading("4.1 The regime flag does not reliably improve forecasts", level=2)
rows = []
for phase in PHASES:
    for model in MODELS:
        for spec in ["M1", "M2"]:
            r = br.real[(br.real.model == model) & (br.real.phase == phase) & (br.real.spec == spec)
                        & (br.real.stress == "all")].iloc[0]
            rows.append([phase, model, spec, f3(r.average_precision), f3(r.roc_auc), f"{r.brier:.4f}"])
table(pd.DataFrame(rows, columns=["Period", "Model", "Spec", "AP (higher better)", "ROC-AUC (higher better)",
                                  "Brier (lower better)"]),
      "Out-of-sample performance without (M1) and with (M2) the regime flag",
      highlight=lambda i, row: row[2] == "M2")
figure(figs["ap"], "AP with and without the regime flag.", "Dotted line = no-skill level.", width=5.8)
rows = []
for phase in PHASES:
    for model in MODELS:
        for metric, lab, d in [("average_precision", "AP", 3), ("roc_auc", "ROC-AUC", 3), ("brier", "Brier", 4)]:
            b, p = br.bt(model, phase, metric), br.ps(model, phase, metric)
            rows.append([phase, model, lab, sgn(b.delta_m2_minus_m1, d),
                         f"[{b.ci_low:+.{d}f}, {b.ci_high:+.{d}f}]", f"{p.placebo_p_value:.2f}"])
table(pd.DataFrame(rows, columns=["Period", "Model", "Metric", "Gain (M2 - M1)", "95% range", "Placebo p"]),
      "Is the gain from the regime flag real?", font=8.5,
      note="For Brier, a negative gain means M2 is better. Shaded rows: placebo p of 0.05 or less.",
      highlight=lambda i, row: float(row[5]) <= 0.05)
para(f"Adding the flag raises AP slightly in every case (Table 3), but none of these gains can be told apart from "
     f"chance (Table 4). Every 95% range for AP includes zero, and 11-22% of fake flags help as much as the real "
     f"one (Figure 4). Only two results pass the 0.05 bar. The forest's 2022 calibration improves (p = "
     f"{br.ps('Random forest', PHASES[0], 'brier').placebo_p_value:.2f}), and the logit's 2023 ROC-AUC rises by "
     f"{d_('Logit', PHASES[1], 'roc_auc'):+.3f} (p = {br.ps('Logit', PHASES[1], 'roc_auc').placebo_p_value:.2f}). "
     f"With twelve comparisons, one or two such results are expected by chance. Meanwhile, 2023 calibration gets "
     f"worse with the flag in both models.")
figure(figs["placebo"], "Real regime flag (coloured line) against 89 fake flags (grey).",
       "A genuinely informative flag would sit to the right of almost all grey bars.")

doc.add_heading("4.2 Why: the flag was right, but rarely switched on", level=2)
rows = []
for phase in PHASES:
    for model in MODELS:
        for s in ["0", "1"]:
            q = br.real[(br.real.model == model) & (br.real.phase == phase) & (br.real.stress == s)]
            if q.empty:
                continue
            a, b = q[q.spec == "M1"].iloc[0], q[q.spec == "M2"].iloc[0]
            rows.append([phase, model, "Stressed" if s == "1" else "Normal", int(a.n), int(a.positives),
                         f3(a.average_precision), f3(b.average_precision), f"{a.brier:.4f}", f"{b.brier:.4f}"])
table(pd.DataFrame(rows, columns=["Period", "Model", "Regime", "Coin-days", "Breaches", "AP M1", "AP M2",
                                  "Brier M1", "Brier M2"]),
      "Performance split by the regime on the forecast day", font=8.5,
      note="AP cannot be computed when every day in a group is a breach (2023 stressed days).")
lg = br.preds[("Logit", PHASES[1])].query("stress == 0")
para(f"In 2022, {br22.loc[1, 'mean']:.0%} of stressed coin-days were breaches against {br22.loc[0, 'mean']:.0%} of "
     f"normal ones, so M2 learned to raise risk under stress and lower it otherwise. In 2023 the flag was still "
     f"right when on: all {int(br23.loc[1, 'size'])} stressed coin-days were breaches. But it was on for only three "
     f"days, and {int(br23.loc[0, 'sum'])} of the {int(p23[TARGET].sum())} breaches came on days it called normal. On "
     f"those days M2 cut its forecasts (logit average {lg.p_m1_oos.mean():.1%} to {lg.p_m2_oos.mean():.1%}), "
     f"which is why calibration worsens (Table 5).")
para("The regime captures crypto-native stress (falling BTC/ETH, fearful sentiment), the setting of the Terra "
     "collapse. SVB was a traditional banking shock: USDC depegged because part of its reserves sat in a failing "
     "bank, while crypto markets were not in a downturn. A crypto-market regime cannot anticipate such a shock.",
     bold_lead="Interpretation. ")

doc.add_heading("4.3 The SVB episode", level=2)
para("Neither model gives USDC or DAI a high depeg probability during SVB (Figure 5). The forest peaks near 25% for "
     "USDC; the logit stays below 7%. The flag did switch on on 10 March, in time for the first USDC and DAI breach "
     "day, but it added only about 3 percentage points to their forecasts (Appendix Table A2). The regime "
     "responded, but the models had learned to give it too little weight for this kind of crisis.")
figure(figs["svb"], "Forecasts for USDC and DAI around the SVB collapse.",
       "Dashed = M1, solid = M2. Triangles = actual breaches. Purple = flag 'stressed'.")

doc.add_heading("4.4 On-chain relationships do change with the regime", level=2)
para(f"Here the answer is yes. In the logit M3, the seven feature-by-stress terms are jointly highly significant "
     f"(Wald chi-square = {wald['wald_chi_square']:.1f}, p = {wald['p_value']:.1e}, date-clustered standard errors). "
     f"In calm markets, the mix of activity matters: more transfers lowers risk, while more unique senders, "
     f"receivers and minting raise it. Under stress these effects fade and volume takes over: higher on-chain "
     f"volume raises risk and more large-transfer volume lowers it (Figure 6, Table 6).")
figure(figs["slopes"], "Logit M3: each feature's effect in normal vs stressed markets (95% intervals).", width=5.8)
s = br.slopes.copy()
table(pd.DataFrame({
    "Feature": s.feature.str.replace("log1p_", "").str.replace("_", " "),
    "Normal market": [f"{a:.2f} ({b:.2f})" for a, b in zip(s.normal_log_odds_slope, s.normal_clustered_se)],
    "Stressed market": [f"{a:.2f} ({b:.2f})" for a, b in zip(s.stressed_log_odds_slope, s.stressed_clustered_se)],
    "Change under stress": s.interaction_delta.map(lambda v: f"{v:+.2f}"),
}), "Logit M3: effect on depeg log-odds per 1 SD of each feature (standard error)", font=8.5,
    note="Transfers, senders and receivers move closely together, so read their effects as a group.")
para(f"The forest finds the same pattern without being told to look for it. Its predicted-risk curves for each "
     f"feature are not parallel between regimes (Figure 7), and they diverge more under the real flag than under "
     f"fake ones (placebo p = {rf_int['joint_placebo_p_value']:.3f}). The difference is largest for burned, minted and "
     f"unique receivers (Table 7).")
figure(figs["pdp"], "Random forest: predicted depeg risk as each feature varies, normal vs stressed (2022).")
s2 = br.rf_int_sum[br.rf_int_sum.statistic != "oob_ap_gain"].copy()
s2["statistic"] = s2.statistic.map(lambda v: "All features combined" if v == "joint_interaction"
                                   else v.replace("I_log1p_", "").replace("_", " "))
table(pd.DataFrame({"Feature": s2.statistic, "Real flag": s2.real.map(lambda v: f"{v:.4f}"),
                    "Average fake flag": s2.placebo_mean.map(lambda v: f"{v:.4f}"),
                    "Placebo p": s2.placebo_p_value.map(lambda v: f"{v:.3f}")}),
      "Random forest: how much each feature's effect depends on the regime", font=8.5,
      note="0 = the regime only shifts risk up or down; larger = the feature's effect itself changes. "
           "Shaded rows: placebo p of 0.05 or less.",
      highlight=lambda i, row: float(row[3]) <= 0.05)
para("The same on-chain movement does mean different things in different markets: heavy volume looks like a run "
     "under stress but is ordinary in calm times. This supports our hypothesis and the regime-dependent logic of "
     "currency-crisis early warning (Kaminsky, Lizondo & Reinhart, 1998). But because the pattern was learned "
     "from Terra, it did not improve forecasts of the next, different crisis.", bold_lead="Interpretation. ")

doc.add_heading("4.5 Early warning: alerts and new episodes", level=2)
rows = []
for r in thr[thr.phase == "2023 frozen test"].itertuples():
    rows.append([r.model, r.spec, f"{r.threshold:.2f}", f3(r.precision), f3(r.recall),
                 r.true_alarms, r.false_alarms, r.missed_breaches])
table(pd.DataFrame(rows, columns=["Model", "Spec", "Cut-off", "Precision", "Recall", "Correct alerts",
                                  "False alarms", "Missed"]),
      "2023 alert performance, cut-offs fixed on 2022 validation", font=8.5,
      note="Precision = share of alerts that were breaches; recall = share of 49 breaches that got an alert. "
           "2022 results are in Appendix Table A3.")
para("Cut-offs set on 2022 (about 0.5-0.6) were far above the probabilities the models produced in 2023, so recall "
     "fell to 2-20% (Table 8). The flag made this worse for the forest (10 breaches caught without it, 3 with it) "
     "by lowering risk on 'normal' days. In 2022 the flag had looked useful, cutting the forest's false alarms "
     "from 8 to 2.")
rows = []
for r in ev[ev.spec == "M2"].pivot_table(index=["coin", "onset_event_date", "stress_on_forecast_day"],
                                         columns="model", values="p_on_onset_forecast").reset_index().itertuples(index=False):
    rows.append([r[0], pd.Timestamp(r[1]).strftime("%d %b %Y"), "Stressed" if r[2] == 1 else "Normal",
                 f"{r[3]:.1%}", f"{r[4]:.1%}"])
table(pd.DataFrame(rows, columns=["Coin", "First breach day", "Regime flag", "Logit M2 forecast", "RF M2 forecast"]),
      "Every new depeg episode in 2023 and the forecast made the day before", font=8.5,
      note="No model (M1 or M2) issued an alert for any of these days or the three days before them.")
para("None of the 10 new episodes in 2023 was warned of by any model, so lead time cannot be measured "
     "(Table 9). Every correct alert fell on a day after an episode had already begun. The models mainly learn "
     "that depegs persist, not when they start, partly because the label rewards spotting ongoing episodes.",
     bold_lead="Interpretation. ")

doc.add_heading("4.6 Robustness", level=2)
lagv = {}
for model, lp in br.lag_preds.items():
    for ph_key, ph in [("2022_validation", PHASES[0]), ("2023_event_study", PHASES[1])]:
        for lag in [0, 1, 3]:
            g = lp[(lp.phase == ph_key) & (lp.regime_lag_days == lag)]
            y = g[TARGET].to_numpy()
            lagv[(model, ph, lag)] = br.score(y, g.p_m2_oos)[0] - br.score(y, g.p_m1_oos)[0]
sc = score_res[score_res.metric == "average_precision"].set_index(["model", "phase"])
cols = [(m, ph) for m in MODELS for ph in PHASES]
rows = []
for lab, fn in [("Main: flag from day t", lambda m, ph: lagv[(m, ph, 0)]),
                ("Flag from day t-1", lambda m, ph: lagv[(m, ph, 1)]),
                ("Flag from day t-3", lambda m, ph: lagv[(m, ph, 3)]),
                ("Stress score (0-3) instead of 0/1 flag", lambda m, ph: sc.loc[(m, ph), "delta"])]:
    rows.append([lab] + [sgn(fn(m, ph)) for m, ph in cols])
table(pd.DataFrame(rows, columns=["Variant", "Logit 2022", "Logit 2023", "RF 2022", "RF 2023"]),
      "AP gain from regime information (M2 - M1) under alternative specifications", font=8.5,
      note="Placebo p-values were computed for the main flag and the stress score; none of their AP gains has a "
           "placebo p below 0.10. Full stress-score results are in Appendix Table A4.")
bullet(" Using the flag from one or three days earlier leaves the small 2023 gains unchanged, but the 2022 "
       "gain swings between +0.04 and +0.20. A genuine signal would not be this sensitive to timing.",
       "Flag timing.")
bullet(" Using the number of stress signals (0-3) instead of the on/off flag does not help, and it removes "
       "the forest's one significant 2022 result.", "Finer regime measure.")
bullet(" The forest is better in 2022 (AP 0.57 vs 0.30 for M1) and the logit in 2023 (0.56 vs 0.51 for M2). "
       "The flexible forest fits Terra-era patterns closely and carries over less well. Neither finds "
       "dependable value in the regime, so the conclusion does not depend on the model.", "Model choice.")

# ============================================================ 5 CONTRIBUTIONS
doc.add_heading("5. Contributions, Limitations, and Future Directions", level=1)
doc.add_heading("5.1 Contributions", level=2)
bullet(" A crypto-market regime does not reliably improve out-of-sample depeg forecasts, in linear or flexible "
       "models. This is the informative null our plan anticipated: the proxy is too crypto-centric to anticipate "
       "a crisis that starts in traditional finance.", "Prediction:")
bullet(" On-chain signals carry different depeg risk in stressed and calm markets. Volume matters under stress; "
       "the composition of activity matters in calm markets. Pooled models (Lee, Chiu & Hsieh, 2025) average "
       "over these two relationships.", "Relationships:")
bullet(" Forward-only testing across two different crises, paired baseline and regime-aware models, and a "
       "placebo test that separates genuine regime information from the benefit of any persistent on/off "
       "variable. The placebo test can be reused for any proposed stress indicator.", "Method:")
bullet(" Good ranking scores (ROC-AUC up to 0.85) coexisted with zero warnings for new episodes. Early-warning "
       "studies should report event-level results, not only AP or ROC-AUC.", "Caution:")
para("For DeFi protocols and exchanges, a crypto stress flag is a useful lens for reading on-chain signals, "
     "but not a substitute for monitoring each stablecoin's reserves and banking exposure, where shocks like SVB "
     "originate.", bold_lead="Practical implication. ")

doc.add_heading("5.2 Limitations", level=2)
bullet(" 3 coins, 8 breaches in the 2022 validation months (all in July), 10 new episodes in 2023.", "Few events:")
bullet(" Counts every breach day, not only new breaches, which pushes models towards detecting persistence.",
       "Label:")
bullet(" The flag is on for only 3 days in 2023, so its test rests largely on the SVB week.",
       "Little regime variation:")
bullet(" 2022 and 2023 on-chain data come from different providers, so part of the gap between years may reflect "
       "measurement differences.", "Two data sources:")
bullet(" Settings were fixed rather than tuned, and data publication timing is assumed rather than verified.",
       "Assumptions:")

doc.add_heading("5.3 Future directions", level=2)
bullet(" Predict the first day of each episode and measure lead time directly.", "Target new episodes:")
bullet(" Add traditional-finance stress (VIX, bank stocks, T-bill spreads) and stablecoin-specific stress "
       "(redemption flows, liquidity-pool imbalances), possibly via a probabilistic regime model. Fix any new "
       "indicator on 2022 only and run it through the same placebo test.", "Better stress measures:")
bullet(" More stablecoins and years would allow testing by holding out one crisis at a time.", "Wider sample:")
bullet(" Re-set alert cut-offs as base rates change, since 2022 cut-offs were far too high for 2023.",
       "Adaptive alerts:")

# ============================================================ appendix
doc.add_page_break()
doc.add_heading("Appendix: supporting tables and figures", level=1)
counters.update({"table": 0, "figure": 0, "prefix": "A"})

rows = []
for model in MODELS:
    pv = br.preds[(model, PHASES[0])]
    for fold, g in pv.groupby("fold"):
        y = g[TARGET].to_numpy()
        a1, _, b1 = br.score(y, g.p_m1_oos)
        a2, _, b2 = br.score(y, g.p_m2_oos)
        rows.append([model, fold, len(g), int(y.sum()), f"{g.stress.mean():.0%}", f3(a1), f3(a2),
                     f"{b1:.4f}", f"{b2:.4f}"])
table(pd.DataFrame(rows, columns=["Model", "Month", "Coin-days", "Breaches", "Stressed", "AP M1", "AP M2",
                                  "Brier M1", "Brier M2"]),
      "2022 validation, month by month", font=8.5,
      note="AP is undefined in months without breaches. The flag improves Brier in July-August but worsens it in "
           "September (mostly stressed, no breaches).")
rows = []
for model in MODELS:
    for phase in PHASES:
        p = br.preds[(model, phase)]
        contrast = p.p_m2_stress1 - p.p_m2_stress0
        contrib = p.p_m2_oos - p.p_m2_stress0
        svb = p[p.coin.isin(["USDC", "DAI"]) & p.event_date.between("2023-03-01", "2023-03-20") & p.stress.eq(1)]
        rows.append([model, phase, f"{100 * contrast.mean():+.1f}", f"{100 * contrib[p.stress.eq(1)].mean():+.1f}",
                     f"{100 * (svb.p_m2_oos - svb.p_m2_stress0).mean():+.1f}" if len(svb) else "-"])
table(pd.DataFrame(rows, columns=["Model", "Period", "Effect of switching flag on (pp)",
                                  "Contribution on stressed days (pp)", "USDC/DAI during SVB (pp)"]),
      "How much of each M2 forecast comes from the regime flag", font=8.5,
      note="pp = percentage points. Same fitted M2, flag forced to 0 vs 1; a model-based breakdown, not causal.")
figure(figs["contrast"], "What a 'stressed' flag would add to the frozen M2 forecast, SVB window.", width=5.8)
rows = []
for r in thr.itertuples():
    rows.append([r.phase, r.model, r.spec, f"{r.threshold:.2f}", f3(r.precision), f3(r.recall), f3(r.f1),
                 r.true_alarms, r.false_alarms, r.missed_breaches])
table(pd.DataFrame(rows, columns=["Period", "Model", "Spec", "Cut-off", "Precision", "Recall", "F1",
                                  "Correct", "False alarms", "Missed"]),
      "Alert performance in both periods", font=8,
      note="The 2022 rows are where the cut-off was chosen, so they flatter the models.")
rows = []
for r in score_res.itertuples():
    lab = {"average_precision": "AP", "roc_auc": "ROC-AUC", "brier": "Brier"}[r.metric]
    d = 4 if r.metric == "brier" else 3
    rows.append([r.phase, r.model, lab, f"{r.M1:.{d}f}", f"{r.M2_flag:.{d}f}", f"{r.M2_score:.{d}f}",
                 f"{r.flag_placebo_p:.2f}", f"{r.placebo_p:.2f}"])
table(pd.DataFrame(rows, columns=["Period", "Model", "Metric", "M1", "M2 flag", "M2 score", "Placebo p (flag)",
                                  "Placebo p (score)"]),
      "Stress score (0-3) vs on/off flag, all metrics", font=8)
figure(figs["coef"], "What each model relies on (trained on all of 2022). Lighter = M1, darker = M2.")
for path, cap in [
    (ROOT / "logit/outputs/logit_decomposition/USDC_M2_decomposition.png",
     "Logit, USDC: M2 forecast with and without the flag's contribution."),
    (ROOT / "logit/outputs/logit_decomposition/DAI_M2_decomposition.png", "Logit, DAI: the same breakdown."),
    (ROOT / "random_forest/outputs/rf_decomposition/USDC_M2_decomposition.png", "Random forest, USDC: M2 breakdown."),
    (ROOT / "random_forest/outputs/rf_decomposition/DAI_M2_decomposition.png", "Random forest, DAI: M2 breakdown."),
    (ROOT / "logit/outputs/regime_lag_sensitivity/USDC_SVB_lag_sensitivity.png", "Logit, USDC: flag from t, t-1, t-3."),
    (ROOT / "logit/outputs/regime_lag_sensitivity/DAI_SVB_lag_sensitivity.png", "Logit, DAI: flag timing."),
    (ROOT / "random_forest/outputs/regime_lag_sensitivity/USDC_SVB_lag_sensitivity.png",
     "Random forest, USDC: flag timing."),
    (ROOT / "random_forest/outputs/regime_lag_sensitivity/DAI_SVB_lag_sensitivity.png", "Random forest, DAI: flag timing."),
]:
    if path.exists():
        figure(path, cap, width=5.5)

doc.save(OUT)
print("Wrote", OUT)


